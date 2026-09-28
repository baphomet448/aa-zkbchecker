"""Excel export for ZKB Checker results."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADERS = [
    "Name", "Found", "Corporation", "Alliance", "Birthday",
    "Kills Total", "Losses Total", "Solo Kills", "Solo Losses",
    "Danger Ratio", "Gang Ratio", "Last Kill Date", "Last Loss Date",
    "Top Ships (Kills)", "Top Ships (Losses)", "Triggers", "Suspicious Alliance",
]

TRIGGER_ROW_FILL = PatternFill(start_color="FF895C", end_color="FF895C", fill_type="solid")
SUSPICIOUS_CELL_FILL = PatternFill(start_color="FF2E2E", end_color="FF2E2E", fill_type="solid")
HEADER_FONT = Font(bold=True)


def _format_ships(ships: list[dict]) -> str:
    if not ships:
        return ""
    return ", ".join(f"{s['name']} ({s['count']})" for s in ships)


def build_workbook(results: list[dict]) -> BytesIO:
    """Build an Excel workbook from a list of character check results
    and return it as an in-memory file-like object, ready to be sent
    as an HTTP response."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT

    for row_index, r in enumerate(results, start=2):
        triggers = r.get("triggers") or []
        suspicious_alliance = r.get("suspicious_alliance") or ""

        values = [
            r.get("name", ""),
            r.get("found", False),
            r.get("corporation_name", ""),
            r.get("alliance_name", ""),
            r.get("birthday", ""),
            r.get("kills_total", ""),
            r.get("losses_total", ""),
            r.get("solo_kills", ""),
            r.get("solo_losses", ""),
            r.get("danger_ratio", ""),
            r.get("gang_ratio", ""),
            r.get("last_kill_date", ""),
            r.get("last_loss_date", ""),
            _format_ships(r.get("top_ships_kills") or []),
            _format_ships(r.get("top_ships_losses") or []),
            ", ".join(triggers),
            suspicious_alliance,
        ]

        for col, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=col, value=value)

        if triggers:
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row=row_index, column=col).fill = TRIGGER_ROW_FILL

        if suspicious_alliance:
            ws.cell(row=row_index, column=len(HEADERS)).fill = SUSPICIOUS_CELL_FILL

    for col in range(1, len(HEADERS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer